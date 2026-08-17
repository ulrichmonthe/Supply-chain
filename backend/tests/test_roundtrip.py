"""End-to-end tests: seed, export, re-import, run, compare, export.

The Excel round-trip test is the important one here. A country that cannot get its
model out of the tool and back in has not been handed anything, and the tool has not
met the durability claim it is sold on.
"""

from __future__ import annotations

import pytest
from sqlalchemy import select

from app.engine import kpis as kpi_mod
from app.engine.roadmap import build_roadmap
from app.engine.runner import run_scenario
from app.io import excel_in, excel_out
from app.io.validation import validate_dataset
from app.models import Country, Demand, Edge, Node, Product, Scenario


@pytest.fixture
def country(session):
    return session.scalar(select(Country).where(Country.code == "PNG"))


def _network(session, country):
    return (
        list(session.scalars(select(Node).where(Node.country_id == country.id))),
        list(session.scalars(select(Edge).where(Edge.country_id == country.id))),
        list(session.scalars(select(Product).where(Product.country_id == country.id))),
        list(session.scalars(select(Demand).where(Demand.country_id == country.id))),
    )


def test_seed_builds_a_connected_network(session, country):
    nodes, edges, products, demand = _network(session, country)
    assert len(nodes) > 130
    assert len(products) == 8
    assert demand

    facilities = {n.id for n in nodes if n.level == 3}
    served = {e.to_node_id for e in edges}
    assert facilities - served == set(), "every facility must have at least one inbound lane"

    hubs = [n for n in nodes if n.hub_capable]
    assert len([h for h in hubs if h.operating_status == "operational"]) == 5


def test_every_seeded_coordinate_is_on_land(session, country):
    """The dataset must pass the check it ships with, or the check is theatre."""
    nodes, edges, products, demand = _network(session, country)
    report = _dataset_report(country, nodes, edges, products, demand)
    offshore = [i for i in report.issues if i.code in ("node.offshore", "node.outside_country")]
    assert offshore == [], [i.message for i in offshore]


def test_seeded_network_passes_its_own_validator(session, country):
    nodes, edges, products, demand = _network(session, country)
    report = _dataset_report(country, nodes, edges, products, demand)
    assert not report.blocking, [i.message for i in report.errors]


def _dataset_report(country, nodes, edges, products, demand):
    """Round-trip through the exporter and importer, then validate what comes back."""
    workbook = excel_out.export_network(country, nodes, edges, products, demand)
    parsed = excel_in.parse_workbook(workbook)
    return validate_dataset(
        country_code=country.code,
        bbox=country.config["bbox"],
        nodes=parsed["nodes"],
        edges=parsed["edges"],
        products=parsed["products"],
        demand=parsed["demand"],
    )


def test_excel_round_trip_preserves_the_model(session, country):
    nodes, edges, products, demand = _network(session, country)
    workbook = excel_out.export_network(country, nodes, edges, products, demand)
    parsed = excel_in.parse_workbook(workbook)

    assert len(parsed["nodes"]) == len(nodes)
    assert len(parsed["edges"]) == len(edges)
    assert len(parsed["products"]) == len(products)
    assert len(parsed["demand"]) == len(demand)

    by_code = {n["code"]: n for n in parsed["nodes"]}
    original = next(n for n in nodes if n.code == "PNG-NCD-001")
    restored = by_code["PNG-NCD-001"]
    assert restored["lat"] == pytest.approx(original.lat)
    assert restored["lon"] == pytest.approx(original.lon)
    assert restored["catchment_population"] == pytest.approx(original.catchment_population)
    assert restored["capacity"]["dry_m3"] == pytest.approx(original.capacity["dry_m3"])

    # A scheduled service must survive with its timetable intact -- it is the whole
    # differentiator, and a round trip that quietly drops it is a silent regression.
    scheduled = next(e for e in parsed["edges"] if e["service_frequency"] == "WEEKLY" and e["mode"] == "sea")
    assert scheduled["capacity_per_trip_m3"] > 0
    assert len(scheduled["monthly_access"]) == 12


def test_template_is_readable_by_the_importer():
    parsed = excel_in.parse_workbook(excel_out.build_template())
    assert parsed["missing_sheets"] == []


def _run_named(session, name: str):
    scenario = session.scalar(select(Scenario).where(Scenario.name.like(f"{name}%")))
    assert scenario is not None, name
    return scenario, run_scenario(session, scenario)


def test_baseline_runs_and_reconciles(session, country):
    scenario, result = _run_named(session, "Baseline")
    assert result.status == "ok"
    kpis = result.kpi_set

    assert kpis["fill_rate"] == pytest.approx(1.0, abs=1e-3)
    assert kpis["total_cost"] > 0
    assert kpis["transport_cost"] + kpis["hub_fixed_cost"] == pytest.approx(kpis["total_cost"], rel=1e-6)

    # Per-facility costs must sum back to the total, or the equity panel is telling
    # a different story from the scorecard.
    per_facility = sum(n["cost"] for n in result.per_node_detail)
    assert per_facility == pytest.approx(kpis["total_cost"], rel=1e-3)

    # Every facility's served + unmet must equal its demand.
    for node in result.per_node_detail:
        assert node["served_m3"] + node["unmet_m3"] == pytest.approx(node["demand_m3"], abs=1e-2)


def test_equity_strata_cover_the_whole_population(session, country):
    _, result = _run_named(session, "Baseline")
    strata = result.equity_detail["strata"]
    assert len(strata) == 5

    total_population = sum(s["population"] for s in strata)
    modelled = sum(n["population"] for n in result.per_node_detail)
    assert total_population == pytest.approx(modelled, rel=0.01)

    # Quintiles are population-weighted, so each should hold a broadly similar share.
    shares = [s["population"] / total_population for s in strata]
    assert min(shares) > 0.10

    # And they must be ordered by vulnerability, which is what makes the panel readable.
    means = [s["mean_vulnerability"] for s in strata]
    assert means == sorted(means)


def test_cost_optimisation_saves_money_by_abandoning_the_vulnerable(session, country):
    """The demo moment that wins a UNICEF room, as an assertion."""
    _, baseline = _run_named(session, "Baseline")
    _, optimised = _run_named(session, "Cost optimisation — unconstrained")

    assert optimised.kpi_set["total_cost"] < baseline.kpi_set["total_cost"]
    saving = 1 - optimised.kpi_set["total_cost"] / baseline.kpi_set["total_cost"]
    assert saving > 0.10

    # ...and the saving comes out of the most vulnerable quintile.
    assert optimised.kpi_set["worst_stratum_fill_rate"] < baseline.kpi_set["worst_stratum_fill_rate"]
    assert optimised.kpi_set["equity_gap"] > baseline.kpi_set["equity_gap"]


def test_equity_floor_buys_back_coverage_at_a_price(session, country):
    _, baseline = _run_named(session, "Baseline")
    _, optimised = _run_named(session, "Cost optimisation — unconstrained")
    _, floored = _run_named(session, "Cost optimisation with a 90% equity floor")

    assert floored.kpi_set["worst_stratum_fill_rate"] >= 0.90 - 1e-6
    assert optimised.kpi_set["total_cost"] < floored.kpi_set["total_cost"] < baseline.kpi_set["total_cost"]


def test_wet_season_costs_more_to_hold_the_same_service(session, country):
    _, baseline = _run_named(session, "Baseline")
    _, march = _run_named(session, "Wet season stress test")

    assert march.kpi_set["fill_rate"] == pytest.approx(baseline.kpi_set["fill_rate"], abs=1e-3)
    assert march.kpi_set["total_cost"] > baseline.kpi_set["total_cost"]
    assert march.solver_log["month"] == 3
    assert march.solver_log["lanes_dropped_by_season"] > 0


def test_thinner_timetables_raise_stockout_risk(session, country):
    """The boat-frequency demo, end to end."""
    _, baseline = _run_named(session, "Baseline")
    _, halved = _run_named(session, "Island services halved")

    assert halved.kpi_set["total_cost"] < baseline.kpi_set["total_cost"]
    assert halved.kpi_set["mean_stockout_risk"] > baseline.kpi_set["mean_stockout_risk"] * 1.4
    assert halved.kpi_set["facilities_at_risk"] > baseline.kpi_set["facilities_at_risk"]

    island = [
        n
        for n in halved.per_node_detail
        if n["terrain_class"] == "island" and n["service_frequency"] == "MONTHLY"
    ]
    assert island, "the frequency override should have reached island services"
    assert max(n["stockout_risk"] for n in island) > 0.5


def test_roadmap_prices_the_change_and_keeps_capex_out_of_opex(session, country):
    baseline_scenario, baseline = _run_named(session, "Baseline")
    scenario, result = _run_named(session, "Open Alotau")

    nodes_by_code = {n.code: n for n in session.scalars(select(Node).where(Node.country_id == country.id))}
    plan = build_roadmap(
        baseline_scenario=baseline_scenario,
        baseline_result=baseline,
        scenario=scenario,
        result=result,
        nodes_by_code=nodes_by_code,
        currency=country.currency,
    )

    titles = " ".join(s["title"] for s in plan["steps"])
    assert "Alotau" in titles and "Wewak" in titles and "Buka" in titles

    # The full capital cost appears once, in the roadmap.
    assert plan["summary"]["one_off_cost_total"] == pytest.approx(4_500_000, rel=0.01)

    # And it does not appear in the annual comparison: only a tenth of it, amortised.
    annual_delta = result.kpi_set["total_cost"] - baseline.kpi_set["total_cost"]
    assert annual_delta < 3_000_000

    for step in plan["steps"]:
        assert step["phase"] in (1, 2, 3)
        assert step["owner"] and step["risk"] and step["evidence"]


def test_scorecard_comparison_labels_direction_correctly(session, country):
    _, baseline = _run_named(session, "Baseline")
    _, optimised = _run_named(session, "Cost optimisation — unconstrained")
    comparison = kpi_mod.compare(baseline.kpi_set, optimised.kpi_set)

    assert comparison["total_cost"]["direction"] == "better"  # cheaper
    assert comparison["worst_stratum_fill_rate"]["direction"] == "worse"  # less equitable
    assert comparison["total_cost"]["delta_pct"] < 0


def test_results_export_contains_the_deliverable_sheets(session, country):
    baseline_scenario, baseline = _run_named(session, "Baseline")
    scenario, result = _run_named(session, "Cost optimisation with a 90% equity floor")

    nodes_by_code = {n.code: n for n in session.scalars(select(Node).where(Node.country_id == country.id))}
    plan = build_roadmap(
        baseline_scenario=baseline_scenario,
        baseline_result=baseline,
        scenario=scenario,
        result=result,
        nodes_by_code=nodes_by_code,
        currency=country.currency,
    )
    payload = excel_out.export_results(
        country, scenario, result, plan, kpi_mod.compare(baseline.kpi_set, result.kpi_set)
    )

    from openpyxl import load_workbook
    import io

    workbook = load_workbook(io.BytesIO(payload))
    assert {"Read Me", "Scorecard", "Facilities", "Lane flows", "Equity", "Roadmap", "Assumptions"} <= set(
        workbook.sheetnames
    )
