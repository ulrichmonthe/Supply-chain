"""An import is a diff with a memory, not a wipe.

The property under test is the one that decides whether anyone will ever correct a
number in the tool: a value typed here has to survive the next master-list refresh.
Around it -- retiring instead of deleting, coming back instead of duplicating, and
the file being allowed to win when a person says so.
"""

from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.db import get_session
from app.io.schema_spec import SHEETS
from app.main import app

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
def client(isolated_session_factory):
    def _session_override():
        session = isolated_session_factory()
        try:
            yield session
        finally:
            session.close()

    app.dependency_overrides[get_session] = _session_override
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()


HUB = ["HUB-1", "Central store", 0, "national_store", -9.4747, 147.1925, "mfl", 0.9,
       "Central", "Port Moresby", "mainland_road", 120000, "operational",
       900, 220, 40, 0, "TRUE", 400000, 0, 9000]


def facility(code="FAC-1", name="Kimbe clinic", population=24000, lat=-5.5510, lon=150.1430):
    return [code, name, 3, "health_centre", lat, lon, "gps", 0.95,
            "West New Britain", "Talasea", "island", population, "operational",
            30, 6, 0, 0, "FALSE", 0, 0, 0]


def workbook(*facilities, demand_quantity=900) -> bytes:
    rows = {
        "Nodes": [HUB, *facilities],
        "Products": [["KIT", "Essential medicines kit", "ambient", 62000.0, 180.0, 720]],
        "Demand": [[f[0], "KIT", 0, demand_quantity, "forecast", 0.6] for f in facilities],
        "Edges": [
            [f"L-{f[0]}", "HUB-1", f[0], "sea", f"Run to {f[0]}", "FORTNIGHTLY", "TUE",
             60, 8, 9000, 0, 0, "", "", "", "", "", "island_sea"] + [1] * 12 + [0.9, 1.5, "TRUE"]
            for f in facilities
        ],
    }
    book = Workbook()
    book.remove(book.active)
    for name, columns in SHEETS.items():
        sheet = book.create_sheet(name)
        sheet.append([column for column, _ in columns])
        for row in rows.get(name, []):
            sheet.append(list(row) + [None] * (len(columns) - len(row)))
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def upload(client: TestClient, content: bytes) -> int:
    report = client.post("/api/countries/1/validate", files={"file": ("net.xlsx", content, XLSX)}).json()
    assert report["blocking"] is False, report["issues"]
    return report["batch_id"]


def preview(client: TestClient, batch_id: int, replace=True) -> dict:
    return client.get(f"/api/imports/{batch_id}/changes", params={"replace": replace}).json()


def commit(client: TestClient, batch_id: int, replace=True, **params):
    response = client.post(f"/api/imports/{batch_id}/commit", params={"replace": replace, **params}, headers={"X-Author": "Importer"})
    assert response.status_code == 200, response.text
    return response.json()


def node(client: TestClient, code: str) -> dict:
    return next(n for n in client.get("/api/countries/1/nodes").json() if n["code"] == code)


# --- the round trip is a no-op ---------------------------------------------------------


def test_reimporting_the_tools_own_export_changes_nothing(client: TestClient):
    """The export and the template share columns so the model survives the tool. The
    diff is what makes that claim checkable: the file the tool wrote must read back as
    no change at all -- not 135 'updates' from a dict shaped differently, and not the
    national store demoted to a clinic by a parser that treated level 0 as missing."""
    exported = client.get("/api/countries/1/export/network.xlsx").content
    batch = upload(client, exported)

    changes = preview(client, batch)

    for sheet in ("nodes", "edges", "products", "demand"):
        counts = changes[sheet]["counts"]
        assert counts["add"] == counts["update"] == counts["retire"] == counts["conflict"] == 0, (sheet, counts, changes[sheet]["updates"][:2])
    assert changes["headline"].startswith("The file matches the model")
    store = next(n for n in client.get("/api/countries/1/nodes").json() if n["level"] == 0)
    commit(client, batch)
    assert next(n for n in client.get("/api/countries/1/nodes").json() if n["id"] == store["id"])["level"] == 0


# --- a full refresh retires, it does not delete -----------------------------------------


def test_the_preview_says_what_a_refresh_would_do_before_it_does_it(client: TestClient):
    seeded = len(client.get("/api/countries/1/nodes").json())
    batch = upload(client, workbook(facility()))

    changes = preview(client, batch)

    assert changes["nodes"]["counts"]["add"] == 2
    assert changes["nodes"]["counts"]["retire"] == seeded
    assert changes["conflicts"] == 0
    assert "would be retired" in changes["headline"]
    assert len(client.get("/api/countries/1/nodes").json()) == seeded, "a preview writes nothing"


def test_a_refresh_retires_what_the_file_no_longer_mentions(client: TestClient):
    seeded = client.get("/api/countries/1/nodes").json()
    results_before = client.get("/api/countries/1/scenarios").json()[0]["latest_result_id"]
    batch = upload(client, workbook(facility()))

    outcome = commit(client, batch)

    assert outcome["counts"]["nodes_retired"] == len(seeded)
    live = {n["code"] for n in client.get("/api/countries/1/nodes").json()}
    assert live == {"HUB-1", "FAC-1"}
    retired = client.get("/api/countries/1/nodes/retired").json()
    assert {n["code"] for n in retired} == {n["code"] for n in seeded}
    assert all(n["retired_reason"].startswith("Not in") for n in retired)
    assert {e["code"] for e in client.get("/api/countries/1/edges").json()} == {"L-FAC-1"}, "retired facilities take their lanes with them"
    assert client.get("/api/countries/1/scenarios").json()[0]["latest_result_id"] == results_before, "past results are untouched"


def test_exports_and_the_overview_do_not_see_retired_rows(client: TestClient):
    exported = load_workbook(io.BytesIO(client.get("/api/countries/1/export/network.xlsx").content))
    codes = [row[0] for row in exported["Nodes"].iter_rows(min_row=2, values_only=True) if row[0]]
    assert sorted(codes) == ["FAC-1", "HUB-1"]
    assert client.get("/api/countries/1/overview").json()["counts"]["facilities"] == 1


# --- the property everything hangs on -------------------------------------------------


def test_a_correction_made_in_the_tool_survives_the_next_import(client: TestClient):
    fac = node(client, "FAC-1")
    client.patch(f"/api/nodes/{fac['id']}", json={"catchment_population": 25000, "reason": "Census 2024."}, headers={"X-Author": "Steward"})

    batch = upload(client, workbook(facility(population=24000)))
    changes = preview(client, batch)
    conflict = changes["nodes"]["conflicts"][0]
    assert conflict["key"] == "FAC-1"
    assert conflict["conflicts"]["catchment_population"] == {"model": 25000.0, "file": 24000, "last_import": 24000}

    outcome = commit(client, batch)

    assert outcome["counts"]["conflicts"] == {"total": 1, "kept": 1, "took_file": 0}
    assert node(client, "FAC-1")["catchment_population"] == 25000
    entry = client.get("/api/countries/1/audit", params={"entity_ref": "FAC-1", "limit": 5}).json()
    kept = [e for e in entry if e["field"] == "catchment_population" and "Kept" in e["rationale"]]
    assert kept and "the file said 24000" in kept[0]["rationale"]


def test_the_file_can_be_told_to_win(client: TestClient):
    batch = upload(client, workbook(facility(population=24000)))

    outcome = commit(client, batch, conflicts="take_file")

    assert outcome["counts"]["conflicts"]["took_file"] == 1
    assert node(client, "FAC-1")["catchment_population"] == 24000


def test_one_conflict_can_be_overridden_while_the_rest_are_kept(client: TestClient):
    fac = node(client, "FAC-1")
    client.patch(f"/api/nodes/{fac['id']}", json={"catchment_population": 26000, "name": "Kimbe Clinic (renamed here)"}, headers={"X-Author": "Steward"})
    batch = upload(client, workbook(facility(population=24000)))
    assert preview(client, batch)["nodes"]["counts"]["conflict"] == 1, "one row, two conflicting fields"

    commit(client, batch, take_file=["nodes:FAC-1:name"])

    after = node(client, "FAC-1")
    assert after["name"] == "Kimbe clinic", "the file won the name"
    assert after["catchment_population"] == 26000, "the correction kept the population"


def test_a_value_the_file_moved_with_no_correction_in_between_is_a_plain_update(client: TestClient):
    # Bring the model back to the file's word first.
    commit(client, upload(client, workbook(facility(population=24000))), conflicts="take_file")
    batch = upload(client, workbook(facility(population=30000)))

    changes = preview(client, batch)

    assert changes["nodes"]["counts"] == {"add": 0, "update": 1, "retire": 0, "restore": 0, "conflict": 0, "unchanged": 1}
    commit(client, batch)
    assert node(client, "FAC-1")["catchment_population"] == 30000


# --- coming back ----------------------------------------------------------------------


def test_a_facility_that_returns_is_restored_not_duplicated(client: TestClient):
    commit(client, upload(client, workbook(facility(population=30000), facility("FAC-2", "Bialla clinic", 9000, -5.32, 151.0))))
    fac2_id = node(client, "FAC-2")["id"]
    commit(client, upload(client, workbook(facility(population=30000))))
    assert "FAC-2" not in {n["code"] for n in client.get("/api/countries/1/nodes").json()}

    batch = upload(client, workbook(facility(population=30000), facility("FAC-2", "Bialla clinic", 9000, -5.32, 151.0)))
    assert preview(client, batch)["nodes"]["counts"]["restore"] == 1
    commit(client, batch)

    back = node(client, "FAC-2")
    assert back["id"] == fac2_id, "the same facility, with its history, not a new row"
    assert "L-FAC-2" in {e["code"] for e in client.get("/api/countries/1/edges").json()}


# --- a merge leaves absent rows alone ---------------------------------------------------


def test_a_merge_never_retires(client: TestClient):
    batch = upload(client, workbook(facility(population=31000)))
    changes = preview(client, batch, replace=False)
    assert changes["nodes"]["counts"]["retire"] == 0
    commit(client, batch, replace=False)
    assert {"FAC-1", "FAC-2", "HUB-1"} <= {n["code"] for n in client.get("/api/countries/1/nodes").json()}
    assert node(client, "FAC-1")["catchment_population"] == 31000
