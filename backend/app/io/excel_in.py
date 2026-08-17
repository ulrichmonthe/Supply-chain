"""Workbook ingest.

Reads the four sheets, normalises them into plain dicts, and hands them to the
validator. Nothing is written to the database until the validation report has been
seen and the import explicitly committed -- an import that half-succeeds is worse
than one that fails, because nobody can tell which half is real.
"""

from __future__ import annotations

import io

from openpyxl import load_workbook

from ..engine import seasonality
from .schema_spec import SHEETS, normalise_header

TRUE_VALUES = {"true", "1", "yes", "y", "t"}


def _to_bool(value, default: bool = False) -> bool:
    if value is None or value == "":
        return default
    return str(value).strip().lower() in TRUE_VALUES


def _to_float(value, default=None):
    if value is None or value == "":
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _read_sheet(workbook, name: str) -> list[dict]:
    if name not in workbook.sheetnames:
        return []
    sheet = workbook[name]
    rows = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        return []

    headers = [normalise_header(h) for h in header_row]
    out: list[dict] = []
    for index, values in enumerate(rows, start=2):
        if values is None or all(v is None or v == "" for v in values):
            continue
        record = {headers[i]: values[i] for i in range(min(len(headers), len(values))) if headers[i]}
        record["_row"] = index
        out.append(record)
    return out


def _monthly_vector(record: dict, prefix: str, profile_key: str, profile_source) -> list[float] | None:
    """Read twelve monthly columns, or expand a named profile, or return None."""
    columns = [f"{prefix}{abbr.lower()}" for abbr in seasonality.MONTH_ABBR]
    if any(column in record and record[column] not in (None, "") for column in columns):
        return [float(_to_float(record.get(column), 1.0)) for column in columns]
    profile_name = record.get(profile_key)
    if profile_name:
        return profile_source(str(profile_name).strip())
    return None


def parse_workbook(data: bytes) -> dict:
    """Turn an uploaded workbook into normalised records."""
    workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)

    nodes_raw = _read_sheet(workbook, "Nodes")
    edges_raw = _read_sheet(workbook, "Edges")
    products_raw = _read_sheet(workbook, "Products")
    demand_raw = _read_sheet(workbook, "Demand")

    missing = [name for name in SHEETS if name not in workbook.sheetnames]

    nodes = []
    for record in nodes_raw:
        nodes.append(
            {
                "_row": record["_row"],
                "code": str(record.get("code") or "").strip(),
                "name": str(record.get("name") or "").strip(),
                "level": int(_to_float(record.get("level"), 3) or 3),
                "type": str(record.get("type") or "health_facility").strip(),
                "lat": _to_float(record.get("lat")),
                "lon": _to_float(record.get("lon")),
                "geocode_source": str(record.get("geocode_source") or "import").strip(),
                "geocode_confidence": _to_float(record.get("geocode_confidence"), 0.5),
                "admin1": str(record.get("admin1") or "").strip() or None,
                "admin2": str(record.get("admin2") or "").strip() or None,
                "terrain_class": str(record.get("terrain_class") or "mainland_road").strip(),
                "catchment_population": _to_float(record.get("catchment_population"), 0.0),
                "operating_status": str(record.get("operating_status") or "operational").strip(),
                "capacity": {
                    "dry_m3": _to_float(record.get("dry_m3"), 0.0),
                    "cold_by_band": {
                        "+2-8": _to_float(record.get("cold_2_8_m3"), 0.0),
                        "-20": _to_float(record.get("cold_minus20_m3"), 0.0),
                        "-70": _to_float(record.get("cold_minus70_m3"), 0.0),
                    },
                },
                "hub_capable": _to_bool(record.get("hub_capable")),
                "hub_fixed_cost": _to_float(record.get("hub_fixed_cost"), 0.0),
                "hub_open_capex": _to_float(record.get("hub_open_capex"), 0.0),
                "hub_throughput_m3": _to_float(record.get("hub_throughput_m3"), 0.0),
                "external_ids": {
                    key: str(record.get(key)).strip()
                    for key in ("dhis2_uid", "msupply_id", "openlmis_code", "mfl_code")
                    if record.get(key)
                },
            }
        )

    edges = []
    for record in edges_raw:
        service_days = record.get("service_days")
        access = _monthly_vector(record, "access_", "access_profile", seasonality.profile)
        cost_profile_name = str(record.get("access_profile") or "").strip()
        edges.append(
            {
                "_row": record["_row"],
                "code": str(record.get("code") or "").strip(),
                "from_node": str(record.get("from_node") or "").strip(),
                "to_node": str(record.get("to_node") or "").strip(),
                "mode": str(record.get("mode") or "road").strip().lower(),
                "service_name": str(record.get("service_name") or "").strip() or None,
                "service_frequency": (
                    str(record.get("service_frequency")).strip().upper()
                    if record.get("service_frequency")
                    else None
                ),
                "service_days": (
                    [d.strip().upper() for d in str(service_days).split(",") if d.strip()]
                    if service_days
                    else None
                ),
                "capacity_per_trip_m3": _to_float(record.get("capacity_per_trip_m3"), 0.0),
                "cold_capacity_per_trip_m3": _to_float(record.get("cold_capacity_per_trip_m3"), 0.0),
                "fixed_cost_per_trip": _to_float(record.get("fixed_cost_per_trip"), 0.0),
                "variable_cost_per_km": _to_float(record.get("variable_cost_per_km"), 0.0),
                "cost_per_m3": _to_float(record.get("cost_per_m3"), 0.0),
                "distance_km": _to_float(record.get("distance_km"), 0.0),
                "base_travel_time_hr": _to_float(record.get("base_travel_time_hr"), 0.0),
                "distance_method": str(record.get("distance_method") or "").strip() or None,
                "distance_confidence": _to_float(record.get("distance_confidence")),
                "distance_note": str(record.get("distance_note") or "").strip() or None,
                "monthly_access": access,
                "monthly_cost_multiplier": (
                    seasonality.cost_profile(cost_profile_name) if cost_profile_name else None
                ),
                "reliability": _to_float(record.get("reliability"), 0.9),
                "lead_time_sd_days": _to_float(record.get("lead_time_sd_days"), 2.0),
                "active": _to_bool(record.get("active"), True),
            }
        )

    products = []
    for record in products_raw:
        products.append(
            {
                "_row": record["_row"],
                "sku": str(record.get("sku") or "").strip(),
                "name": str(record.get("name") or "").strip(),
                "temperature_band": str(record.get("temperature_band") or "ambient").strip(),
                "volume_per_unit_cm3": _to_float(record.get("volume_per_unit_cm3")),
                "unit_cost": _to_float(record.get("unit_cost"), 0.0),
                "shelf_life_days": int(_to_float(record.get("shelf_life_days"), 730) or 730),
            }
        )

    demand = []
    for record in demand_raw:
        demand.append(
            {
                "_row": record["_row"],
                "node": str(record.get("node") or "").strip(),
                "product": str(record.get("product") or "").strip(),
                "period": int(_to_float(record.get("period"), 0) or 0),
                "quantity": _to_float(record.get("quantity")),
                "source": str(record.get("source") or "proxy").strip(),
                "confidence": _to_float(record.get("confidence"), 0.5),
            }
        )

    workbook.close()
    return {
        "nodes": nodes,
        "edges": edges,
        "products": products,
        "demand": demand,
        "missing_sheets": missing,
    }
