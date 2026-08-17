"""Workbook export.

Three products come out of this module:

    template     an empty workbook with the column contract and a Read Me sheet
    network      everything currently loaded, in a form that can be edited and
                 uploaded straight back
    results      the scorecard, per-facility detail, equity decomposition and the
                 costed roadmap

The results workbook is the consultancy deliverable. It is deliberately an Excel
file rather than a PDF, because the expert will want to add the two columns the
model does not know about, and a government analyst will want to re-sort it.
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..engine import seasonality
from ..engine.kpis import KPI_META, KPI_ORDER
from .schema_spec import DEMAND_COLUMNS, EDGE_COLUMNS, NODE_COLUMNS, PRODUCT_COLUMNS, SHEETS

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=13)
NOTE_FONT = Font(italic=True, color="555555")


def _write_header(sheet, columns: list[str], row: int = 1) -> None:
    for index, name in enumerate(columns, start=1):
        cell = sheet.cell(row=row, column=index, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = sheet.cell(row=row + 1, column=1)


def _autosize(sheet, max_width: int = 46) -> None:
    widths: dict[int, int] = {}
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            widths[cell.column] = min(max_width, max(widths.get(cell.column, 10), len(str(cell.value)) + 2))
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width


def _readme(workbook: Workbook, title: str, notes: list[str]) -> None:
    sheet = workbook.create_sheet("Read Me", 0)
    sheet["A1"] = title
    sheet["A1"].font = TITLE_FONT
    row = 3
    for note in notes:
        sheet.cell(row=row, column=1, value=note).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[row].height = 30
        row += 1
    sheet.column_dimensions["A"].width = 120


def build_template() -> bytes:
    """A blank workbook that documents itself."""
    workbook = Workbook()
    workbook.remove(workbook.active)

    for sheet_name, columns in SHEETS.items():
        sheet = workbook.create_sheet(sheet_name)
        _write_header(sheet, [name for name, _ in columns])
        for index, (_, description) in enumerate(columns, start=1):
            sheet.cell(row=2, column=index, value=description).font = NOTE_FONT
        sheet.cell(row=2, column=1).comment = None
        _autosize(sheet)

    reference = workbook.create_sheet("Reference")
    reference["A1"] = "Named seasonal access profiles"
    reference["A1"].font = TITLE_FONT
    _write_header(reference, ["profile", *seasonality.MONTH_ABBR], row=3)
    row = 4
    for name, vector in seasonality.PROFILES.items():
        reference.cell(row=row, column=1, value=name)
        for index, value in enumerate(vector, start=2):
            reference.cell(row=row, column=index, value=value)
        row += 1
    _autosize(reference)

    _readme(
        workbook,
        "Health Supply Chain Network Design — data template",
        [
            "Row 2 of each sheet describes the column. Delete row 2 before you upload, or "
            "leave it: the importer skips rows whose code column is not a real code.",
            "Nodes, Products and Demand are required. Edges are required for any facility "
            "you expect the model to reach.",
            "Distances can be left blank. The distance cascade will compute them and record "
            "how: a routed distance from OSRM, or a great-circle distance inflated by a "
            "terrain-specific detour factor. Type a distance yourself and set distance_method "
            "to 'manual' when you have a measured figure — a manual value always wins.",
            "Seasonal access is twelve numbers per lane, January to December, from 0 "
            "(impassable) to 1 (normal). Instead of typing twelve values you can name a "
            "profile in access_profile; the Reference sheet lists them.",
            "A scheduled sea or air service needs service_frequency AND capacity_per_trip_m3. "
            "The timetable is the point: a fortnightly boat with a 12 m³ hold is a completely "
            "different network from a weekly one, and the model will show you the difference.",
            "Mark every demand row as actual, forecast or proxy. Proxy means derived from "
            "population rather than measured. It is a legitimate starting point — but the "
            "tool will say so on every output, and so should you.",
        ],
    )

    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def export_network(country, nodes, edges, products, demand) -> bytes:
    """Everything currently loaded, in upload-ready form."""
    workbook = Workbook()
    workbook.remove(workbook.active)

    node_sheet = workbook.create_sheet("Nodes")
    _write_header(node_sheet, [name for name, _ in NODE_COLUMNS])
    for row_index, node in enumerate(nodes, start=2):
        capacity = node.capacity or {}
        cold = capacity.get("cold_by_band", {}) or {}
        external = node.external_ids or {}
        values = [
            node.code, node.name, node.level, node.type, node.lat, node.lon,
            node.geocode_source, node.geocode_confidence, node.admin1, node.admin2,
            node.terrain_class, node.catchment_population, node.operating_status,
            capacity.get("dry_m3", 0.0), cold.get("+2-8", 0.0), cold.get("-20", 0.0),
            cold.get("-70", 0.0), node.hub_capable, node.hub_fixed_cost, node.hub_open_capex,
            node.hub_throughput_m3, external.get("dhis2_uid"), external.get("msupply_id"),
            external.get("openlmis_code"), external.get("mfl_code"),
        ]
        for column_index, value in enumerate(values, start=1):
            node_sheet.cell(row=row_index, column=column_index, value=value)
    _autosize(node_sheet)

    edge_sheet = workbook.create_sheet("Edges")
    _write_header(edge_sheet, [name for name, _ in EDGE_COLUMNS])
    for row_index, edge in enumerate(edges, start=2):
        access = edge.monthly_access or [1.0] * 12
        values = [
            edge.code, edge.from_node.code, edge.to_node.code, edge.mode, edge.service_name,
            edge.service_frequency,
            ",".join(edge.service_days) if edge.service_days else None,
            edge.capacity_per_trip_m3, edge.cold_capacity_per_trip_m3, edge.fixed_cost_per_trip,
            edge.variable_cost_per_km, edge.cost_per_m3, edge.distance_km,
            edge.base_travel_time_hr, edge.distance_method, edge.distance_confidence,
            edge.distance_note, None, *access, edge.reliability, edge.lead_time_sd_days,
            edge.active,
        ]
        for column_index, value in enumerate(values, start=1):
            edge_sheet.cell(row=row_index, column=column_index, value=value)
    _autosize(edge_sheet)

    product_sheet = workbook.create_sheet("Products")
    _write_header(product_sheet, [name for name, _ in PRODUCT_COLUMNS])
    for row_index, product in enumerate(products, start=2):
        values = [
            product.sku, product.name, product.temperature_band,
            product.volume_per_unit_cm3, product.unit_cost, product.shelf_life_days,
        ]
        for column_index, value in enumerate(values, start=1):
            product_sheet.cell(row=row_index, column=column_index, value=value)
    _autosize(product_sheet)

    node_codes = {n.id: n.code for n in nodes}
    product_codes = {p.id: p.sku for p in products}
    demand_sheet = workbook.create_sheet("Demand")
    _write_header(demand_sheet, [name for name, _ in DEMAND_COLUMNS])
    for row_index, row in enumerate(demand, start=2):
        values = [
            node_codes.get(row.node_id), product_codes.get(row.product_id),
            row.period, row.quantity, row.source, row.confidence,
        ]
        for column_index, value in enumerate(values, start=1):
            demand_sheet.cell(row=row_index, column=column_index, value=value)
    _autosize(demand_sheet)

    _readme(
        workbook,
        f"{country.name} — network export",
        [
            "This is the complete model as currently loaded. Edit it and upload it back "
            "through Data → Import; the column layout is identical to the blank template.",
            "Distances carry the method that produced them. Anything marked detour_factor is "
            "an estimate from a great-circle distance, and is the first thing to replace with "
            "a measured figure when you have one.",
            "Nothing in this file is a secret and nothing needs the application to read it. "
            "That is deliberate: the model has to survive the end of the engagement.",
        ],
    )

    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def export_results(country, scenario, result, roadmap: dict | None, comparison: dict | None) -> bytes:
    """The consultancy deliverable."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    currency = country.currency
    kpis = result.kpi_set or {}
    log = result.solver_log or {}

    summary = workbook.create_sheet("Scorecard")
    summary["A1"] = f"{country.name} — {scenario.name}"
    summary["A1"].font = TITLE_FONT
    summary["A2"] = scenario.description or ""
    summary["A3"] = (
        f"Conditions: {log.get('month_label', 'Annualised')} · "
        f"solver {log.get('solver', 'HiGHS')} {log.get('status', '')} · "
        f"run {result.run_timestamp:%Y-%m-%d %H:%M} UTC · {result.runtime_ms} ms"
    )
    summary["A3"].font = NOTE_FONT

    _write_header(summary, ["Indicator", f"Value ({currency})", "Unit", "vs baseline", "Direction"], row=5)
    row = 6
    for key in KPI_ORDER:
        if key not in kpis:
            continue
        meta = KPI_META[key]
        delta = (comparison or {}).get(key, {})
        summary.cell(row=row, column=1, value=meta["label"])
        summary.cell(row=row, column=2, value=kpis[key])
        summary.cell(row=row, column=3, value=meta["unit"])
        summary.cell(row=row, column=4, value=delta.get("delta"))
        summary.cell(row=row, column=5, value=delta.get("direction"))
        row += 1
    _autosize(summary)

    detail = workbook.create_sheet("Facilities")
    columns = [
        "code", "name", "admin1", "population", "demand_m3", "served_m3", "fill_rate",
        "cost", "cost_per_capita", "primary_mode", "service_name", "service_frequency",
        "stockout_risk", "storage_days", "vulnerability", "stratum", "restricted_months",
        "reachable",
    ]
    _write_header(detail, columns)
    for row_index, node in enumerate(result.per_node_detail or [], start=2):
        for column_index, key in enumerate(columns, start=1):
            detail.cell(row=row_index, column=column_index, value=node.get(key))
    _autosize(detail)

    flows = workbook.create_sheet("Lane flows")
    flow_columns = [
        "edge_code", "hub_code", "facility_code", "mode", "service_name", "frequency",
        "volume_m3", "unit_cost", "cost", "distance_km", "distance_method",
        "distance_confidence", "capacity_m3", "capacity_utilisation",
    ]
    _write_header(flows, flow_columns)
    for row_index, flow in enumerate(result.per_edge_flow or [], start=2):
        for column_index, key in enumerate(flow_columns, start=1):
            flows.cell(row=row_index, column=column_index, value=flow.get(key))
    _autosize(flows)

    equity_sheet = workbook.create_sheet("Equity")
    equity = result.equity_detail or {}
    equity_sheet["A1"] = "Who the plan reaches, by vulnerability quintile"
    equity_sheet["A1"].font = TITLE_FONT
    equity_sheet["A2"] = (
        "Quintiles are population-weighted: each band is a fifth of the people, not a fifth "
        "of the facilities. Vulnerability is structural — remoteness, terrain, months of "
        "restricted access and mode dependency — and is computed before any optimisation, "
        "so it cannot be gamed by the solution."
    )
    equity_sheet["A2"].font = NOTE_FONT
    equity_columns = [
        "label", "facilities", "population", "demand", "served", "fill_rate",
        "cost", "cost_per_capita", "mean_vulnerability",
    ]
    _write_header(equity_sheet, equity_columns, row=4)
    for row_index, bucket in enumerate(equity.get("strata", []), start=5):
        for column_index, key in enumerate(equity_columns, start=1):
            equity_sheet.cell(row=row_index, column=column_index, value=bucket.get(key))
    _autosize(equity_sheet)

    if roadmap:
        sheet = workbook.create_sheet("Roadmap")
        sheet["A1"] = f"Costed implementation roadmap — {roadmap['scenario']} vs {roadmap['baseline']}"
        sheet["A1"].font = TITLE_FONT
        totals = roadmap["summary"]
        sheet["A2"] = (
            f"One-off cost {totals['one_off_cost_total']:,.0f} {currency} · "
            f"annual change {totals['annual_cost_delta']:,.0f} {currency} · "
            f"payback {totals['payback_years'] if totals['payback_years'] else 'n/a'} years"
        )
        sheet["A2"].font = NOTE_FONT
        roadmap_columns = [
            "id", "phase_label", "phase_window", "title", "detail", "one_off_cost",
            "annual_cost_delta", "owner", "risk", "evidence",
        ]
        _write_header(sheet, roadmap_columns, row=4)
        for row_index, step in enumerate(roadmap["steps"], start=5):
            for column_index, key in enumerate(roadmap_columns, start=1):
                cell = sheet.cell(row=row_index, column=column_index, value=step.get(key))
                if key in ("detail", "risk", "evidence"):
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
        _autosize(sheet, max_width=60)

    assumptions = workbook.create_sheet("Assumptions")
    assumptions["A1"] = "Assumptions behind these numbers"
    assumptions["A1"].font = TITLE_FONT
    _write_header(assumptions, ["Item", "Value"], row=3)
    rows = [
        ("Run conditions", log.get("month_label", "Annualised")),
        ("Solver", f"{log.get('solver', 'HiGHS')} — {log.get('status', '')}"),
        ("Facilities in model", log.get("facilities")),
        ("Lanes in model", log.get("lanes")),
        ("Lanes closed by season", log.get("lanes_dropped_by_season")),
        ("Hubs open", ", ".join(log.get("hubs_open_codes", []))),
        ("Objective weights", str(log.get("objective_weights"))),
        ("Penalty per m³ unserved", log.get("base_unmet_penalty_per_m3")),
        ("Demand growth applied", log.get("demand_growth_applied")),
        ("Levers", str(scenario.levers)),
        ("Constraints", str(scenario.constraints)),
    ]
    for row_index, (label, value) in enumerate(rows, start=4):
        assumptions.cell(row=row_index, column=1, value=label)
        assumptions.cell(row=row_index, column=2, value=str(value) if value is not None else "")
    _autosize(assumptions, max_width=90)

    _readme(
        workbook,
        f"{country.name} — {scenario.name} — results",
        [
            "Scorecard is the headline comparison. Facilities is the per-site detail behind "
            "it. Lane flows is what moved where, and at what unit cost.",
            "Equity shows who the plan reaches. Read it next to the Scorecard, never instead "
            "of it: a cost saving that comes out of the most vulnerable quintile is a "
            "transfer, not a saving, and this sheet is where that becomes visible.",
            "Assumptions lists every setting that produced these numbers. If a figure is "
            "challenged in a workshop, start there.",
        ],
    )

    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()
